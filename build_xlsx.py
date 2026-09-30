"""
Генерация XLSX по категориям из dump.json.

Выход:
  - vacancies_vrachi.xlsx            — все «Врач»
  - vacancies_sredniy_medpersonal.xlsx — все «Средний медицинский персонал»
  - vacancies_mladshiy_medpersonal.xlsx — все «Младший медицинский персонал»
  - for_universities.xlsx            — «Врач» + «Средний медицинский персонал»
  - for_colleges.xlsx                — «Средний медицинский персонал» + «Младший медицинский персонал»

Зависимость: pip install openpyxl
"""

import json
import re
import sys
from typing import Callable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

IN_FILE = "dump.json"

# --- идентификаторы категорий из /admin/vacancy/category/index ---
CAT_DOCTOR = "1"   # Врач
CAT_MIDDLE = "2"   # Средний медицинский персонал
CAT_JUNIOR = "3"   # Младший медицинский персонал

COLUMNS = [
    ("title",          "Название вакансии",         50),
    ("speciality",     "Специальность",             28),
    ("organization",   "Организация",               45),
    ("area",           "Район",                     22),
    ("work_mode",      "Режим работы",              22),
    ("org_address",    "Адрес организации",         45),
    ("org_phone",      "Телефон организации",       22),
    ("org_website",    "Сайт",                      32),
    ("hr_contact",     "Контактное лицо (кадры)",   30),
    ("hr_phone",       "Телефон кадров",            22),
    ("hr_email",       "E-mail для вакансий",       28),
    ("social_support", "Меры соцподдержки",         65),
]


# ---------- загрузка ----------
def load(path: str = IN_FILE) -> dict:
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except FileNotFoundError:
        print(f"ERROR: {path} не найден", file=sys.stderr)
        sys.exit(1)
    except json.JSONDecodeError as e:
        print(f"ERROR: {path} не валидный JSON: {e}", file=sys.stderr)
        sys.exit(1)


def id_from_url(url: str) -> str | None:
    m = re.search(r"[?&]id=(\d+)", url or "")
    return m.group(1) if m else None


def field_value(field):
    if isinstance(field, dict):
        return field.get("value")
    return field


def dash(v) -> str:
    if v is None:
        return "—"
    s = str(v).strip()
    return s if s else "—"


# ---------- справочники ----------
def load_categories(data: dict) -> dict[str, dict]:
    result: dict[str, dict] = {}
    for row in (data.get("categories") or {}).get("items") or []:
        cat_id = row.get("id")
        if not isinstance(cat_id, int):
            continue
        form = ((row.get("detail") or {}).get("form") or {}).get("VacancyCategory") or {}
        result[str(cat_id)] = {"name": form.get("name")}
    return result


def load_vacancy_list(data: dict) -> dict[str, dict]:
    result: dict[str, dict] = {}
    for row in (data.get("vacancy_list") or {}).get("items") or []:
        item_id = row.get("id")
        if not isinstance(item_id, int):
            continue
        form = ((row.get("detail") or {}).get("form") or {}).get("VacancyList") or {}
        cat = form.get("category") or {}
        result[str(item_id)] = {
            "name":           form.get("name") or row.get("name"),
            "category_id":    str(cat.get("value") or "") or None,
            "category_label": cat.get("label"),
        }
    return result


def load_organizations(data: dict) -> dict[str, dict]:
    result: dict[str, dict] = {}
    for row in (data.get("organizations") or {}).get("items") or []:
        detail = row.get("detail") or {}
        org_id = id_from_url(detail.get("url") or "")
        if org_id is None:
            v = row.get("id")
            org_id = str(v) if isinstance(v, int) else None
        if org_id is None:
            continue
        form = (detail.get("form") or {}).get("Organization") or {}
        result[org_id] = {
            "name":                  form.get("name"),
            "short_name":            form.get("short_name"),
            "address":               form.get("address"),
            "website":               form.get("website"),
            "phone":                 form.get("phone"),
            "personnel_dep_address": form.get("personnel_dep_address"),
            "personnel_dep_contact": form.get("personnel_dep_contact"),
            "personnel_dep_phone":   form.get("personnel_dep_phone"),
            "personnel_dep_email":   form.get("personnel_dep_email"),
            "territory":             row.get("territory"),
        }
    return result


# ---------- соцподдержка ----------
def format_social(form: dict) -> str:
    parts: list[str] = []

    def yes(key: str) -> bool:
        return str(field_value(form.get(key)) or "").strip() == "1"

    if yes("housing"):
        parts.append("возможность получения служебного жилья")
    if yes("social_land"):
        parts.append("безвозмездное использование земельного участка")
    if yes("social_communal"):
        parts.append("компенсация оплаты ЖКХ")
    if yes("social_rent"):
        parts.append("компенсация аренды жилья")
    if yes("social_mortgage"):
        parts.append("компенсация первоначального взноса по ипотеке")
    if yes("social_deposit"):
        parts.append("социальная выплата по вкладу для накопления средств")
    if yes("zemskii"):
        raw = str(form.get("zemskii_sum") or "").strip().replace(" ", "")
        if raw:
            try:
                n = int(float(raw.replace(",", ".")))
                parts.append(f"выплата по программе «Земский доктор/фельдшер» — {n:,} ₽".replace(",", " "))
            except ValueError:
                parts.append(f"выплата по программе «Земский доктор/фельдшер» — {raw} ₽")
        else:
            parts.append("выплата по программе «Земский доктор/фельдшер»")

    if not parts:
        return "—"
    parts[0] = parts[0][0].upper() + parts[0][1:]
    return "; ".join(parts) + "."


# ---------- сборка строк ----------
def build_rows(data: dict) -> list[dict]:
    cats = load_categories(data)
    vlist = load_vacancy_list(data)
    orgs = load_organizations(data)

    out: list[dict] = []

    for vac in (data.get("vacancies") or {}).get("items") or []:
        detail = vac.get("detail") or {}
        if "error" in detail:
            continue
        form = (detail.get("form") or {}).get("Vacancy") or {}

        vl_id = str(form.get("vacancy_category") or "").strip()
        vl_entry = vlist.get(vl_id) or {}
        cat_id = vl_entry.get("category_id")
        cat_label = vl_entry.get("category_label") or (cats.get(cat_id or "") or {}).get("name")
        if not cat_id:
            continue

        org_id = str(field_value(form.get("organization")) or "")
        org = orgs.get(org_id) or {}

        terr = form.get("territory")
        if isinstance(terr, dict):
            area = terr.get("label") or ""
        elif isinstance(terr, str):
            area = terr
        else:
            area = ""
        if not area:
            area = org.get("territory") or ""

        title = (
            str(form.get("name_alt") or "").strip()
            or str(form.get("name") or "").strip()
            or str(vl_entry.get("name") or "").strip()
        )

        out.append({
            "title":          dash(title),
            "speciality":     dash(cat_label),
            "organization":   dash(org.get("short_name") or org.get("name")),
            "area":           dash(area),
            "work_mode":      dash(form.get("work_mode")),
            "org_address":    dash(org.get("address")),
            "org_phone":      dash(org.get("phone") or org.get("personnel_dep_phone")),
            "org_website":    dash(org.get("website")),
            "hr_contact":     dash(org.get("personnel_dep_contact")),
            "hr_phone":       dash(org.get("personnel_dep_phone")),
            "hr_email":       dash(org.get("personnel_dep_email")),
            "social_support": format_social(form),
            "_cat_id":        cat_id,
        })

    return out


# ---------- запись xlsx ----------
def write_xlsx(rows: list[dict], path: str, sheet_name: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]

    for i, (_, title, width) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=i, value=title)
        c.font = Font(bold=True)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 24

    for r_idx, row in enumerate(rows, start=2):
        for c_idx, (key, _, _) in enumerate(COLUMNS, start=1):
            c = ws.cell(row=r_idx, column=c_idx, value=row.get(key, ""))
            c.alignment = Alignment(vertical="top", wrap_text=True)

    ws.freeze_panes = "A2"
    wb.save(path)


# ---------- цели ----------
TARGETS: list[dict] = [
    {
        "label":     "Врач (все)",
        "sheet":     "Врач",
        "file":      "vacancies_vrachi.xlsx",
        "predicate": lambda r: r["_cat_id"] == CAT_DOCTOR,
    },
    {
        "label":     "Средний медперсонал (все)",
        "sheet":     "Средний медперсонал",
        "file":      "vacancies_sredniy_medpersonal.xlsx",
        "predicate": lambda r: r["_cat_id"] == CAT_MIDDLE,
    },
    {
        "label":     "Младший медперсонал (все)",
        "sheet":     "Младший медперсонал",
        "file":      "vacancies_mladshiy_medpersonal.xlsx",
        "predicate": lambda r: r["_cat_id"] == CAT_JUNIOR,
    },
    {
        "label":     "Для вузов: Врач + Средний медперсонал",
        "sheet":     "Вузы",
        "file":      "for_universities.xlsx",
        "predicate": lambda r: r["_cat_id"] in (CAT_DOCTOR, CAT_MIDDLE),
    },
    {
        "label":     "Для колледжей: Средний + Младший медперсонал",
        "sheet":     "Колледжи",
        "file":      "for_colleges.xlsx",
        "predicate": lambda r: r["_cat_id"] in (CAT_MIDDLE, CAT_JUNIOR),
    },
]


# ---------- main ----------
def main() -> int:
    data = load(IN_FILE)
    all_rows = build_rows(data)

    if not all_rows:
        print("ERROR: не собрано ни одной строки — проверь dump.json", file=sys.stderr)
        return 1

    print(f"Всего вакансий с известной категорией: {len(all_rows)}\n")

    for t in TARGETS:
        pred: Callable[[dict], bool] = t["predicate"]
        rows = [r for r in all_rows if pred(r)]
        if not rows:
            print(f"  [!] {t['label']}: 0 вакансий — пропускаю")
            continue
        write_xlsx(rows, t["file"], sheet_name=t["sheet"])
        print(f"  {t['label']}: {len(rows)} → {t['file']}")

    return 0


if __name__ == "__main__":
    sys.exit(main())